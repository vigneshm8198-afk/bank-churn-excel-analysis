"""Build the Bank Customer Churn Excel workbook (formula-driven) from a CSV.

Usage:
    python src/build_workbook.py                      # uses data/Churn_Modelling.csv, else the sample
    python src/build_workbook.py --input my.csv --output excel/out.xlsx

Every number in the workbook is an Excel formula, so it recalculates if you
edit the Data sheet. Python only lays out the sheets and writes the raw data.
"""
import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
REAL = ROOT / "data" / "Churn_Modelling.csv"
SAMPLE = ROOT / "data" / "sample" / "sample_churn_data.csv"
REQUIRED = ["RowNumber", "CustomerId", "Surname", "CreditScore", "Geography", "Gender", "Age",
            "Tenure", "Balance", "NumOfProducts", "HasCrCard", "IsActiveMember",
            "EstimatedSalary", "Exited"]

FONT = "Arial"
NAVY, BLUE, LIGHT, GREY = "1F3864", "2E75B6", "DDEBF7", "F2F2F2"
f_title = Font(name=FONT, size=18, bold=True, color=NAVY)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_body = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_note = Font(name=FONT, size=9, italic=True, color="595959")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_light = PatternFill("solid", fgColor=LIGHT)
fill_grey = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

# Helper-column rules (documented on the README sheet). {r} = row number.
HELPERS = [
    ("Age Group", '=IF({Age}{r}<30,"18 to 29",IF({Age}{r}<40,"30 to 39",IF({Age}{r}<50,"40 to 49",IF({Age}{r}<60,"50 to 59","60 plus"))))',
     ["18 to 29", "30 to 39", "40 to 49", "50 to 59", "60 plus"]),
    ("Credit Band", '=IF({CreditScore}{r}<580,"Poor (300-579)",IF({CreditScore}{r}<670,"Fair (580-669)",IF({CreditScore}{r}<740,"Good (670-739)",IF({CreditScore}{r}<800,"Very Good (740-799)","Excellent (800-850)"))))',
     ["Poor (300-579)", "Fair (580-669)", "Good (670-739)", "Very Good (740-799)", "Excellent (800-850)"]),
    ("Balance Segment", '=IF({Balance}{r}=0,"Zero",IF({Balance}{r}<=100000,"Up to 100K",IF({Balance}{r}<=150000,"100K to 150K","Above 150K")))',
     ["Zero", "Up to 100K", "100K to 150K", "Above 150K"]),
    ("Tenure Band", '=IF({Tenure}{r}<=2,"0 to 2 yrs",IF({Tenure}{r}<=5,"3 to 5 yrs",IF({Tenure}{r}<=8,"6 to 8 yrs","9 to 10 yrs")))',
     ["0 to 2 yrs", "3 to 5 yrs", "6 to 8 yrs", "9 to 10 yrs"]),
    ("Activity Status", '=IF({IsActiveMember}{r}=1,"Active","Inactive")', ["Active", "Inactive"]),
]


def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = f_head, fill_head, box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def load(path):
    df = pd.read_csv(path)
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        sys.exit(f"ERROR: {path} is missing columns: {missing}")
    return df[REQUIRED].copy()


def build(df, out, source_label):
    n = len(df)
    last = n + 1
    wb = Workbook()
    ws_readme, ws_dash = wb.active, wb.create_sheet("Dashboard")
    ws_readme.title = "README"
    ws_an, ws_data = wb.create_sheet("Analysis"), wb.create_sheet("Data")

    # ---------------- Data sheet ----------------
    letters = {name: get_column_letter(i + 1) for i, name in enumerate(REQUIRED)}
    hnames = [h[0] for h in HELPERS]
    for j, h in enumerate(hnames):
        letters[h] = get_column_letter(len(REQUIRED) + 1 + j)
    for j, name in enumerate(REQUIRED + hnames, start=1):
        ws_data.cell(row=1, column=j, value=name)
    style_header(ws_data, 1, 1, len(REQUIRED) + len(hnames))
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            ws_data.cell(row=i, column=j, value=v.item() if hasattr(v, "item") else v)
        for j, (name, tmpl, _) in enumerate(HELPERS):
            ws_data.cell(row=i, column=len(REQUIRED) + 1 + j,
                         value=tmpl.format(r=i, **{k: letters[k] for k in REQUIRED}))
    for j in range(1, len(REQUIRED) + len(hnames) + 1):
        ws_data.column_dimensions[get_column_letter(j)].width = 16
    ws_data.freeze_panes = "A2"
    ws_data.auto_filter.ref = f"A1:{get_column_letter(len(REQUIRED) + len(hnames))}{last}"
    for col in ("Balance", "EstimatedSalary"):
        for r in range(2, last + 1):
            ws_data[f"{letters[col]}{r}"].number_format = "#,##0.00"

    def rng(name):
        return f"Data!${letters[name]}$2:${letters[name]}${last}"

    EX = rng("Exited")

    # ---------------- Analysis sheet ----------------
    ws_an["A1"] = "Churn Analysis Tables (all values are formulas on the Data sheet)"
    ws_an["A1"].font = f_title
    dims = [("Geography", sorted(df["Geography"].unique())),
            ("Gender", sorted(df["Gender"].unique()))]
    dims += [(h[0], h[2]) for h in HELPERS]
    dims.insert(4, ("NumOfProducts", sorted(df["NumOfProducts"].unique().tolist())))
    dims.insert(2, ("HasCrCard", None))  # placeholder removed below
    dims = [d for d in dims if d[1] is not None]

    tables = {}
    row = 3
    for name, cats in dims:
        ws_an.cell(row=row, column=1, value=f"Churn by {name}").font = Font(name=FONT, size=12, bold=True, color=NAVY)
        hdr = row + 1
        for c, t in enumerate([name, "Customers", "Churned", "Churn Rate", "Share of All Churners", "Rate vs Overall (pts)"], 1):
            ws_an.cell(row=hdr, column=c, value=t)
        style_header(ws_an, hdr, 1, 6)
        first = hdr + 1
        for k, cat in enumerate(cats):
            r = first + k
            ws_an.cell(row=r, column=1, value=cat if isinstance(cat, str) else int(cat))
            ws_an.cell(row=r, column=2, value=f"=COUNTIFS({rng(name)},$A{r})")
            ws_an.cell(row=r, column=3, value=f"=COUNTIFS({rng(name)},$A{r},{EX},1)")
            ws_an.cell(row=r, column=4, value=f"=IFERROR(C{r}/B{r},0)")
        lastcat = first + len(cats) - 1
        tot = lastcat + 1
        ws_an.cell(row=tot, column=1, value="Total")
        ws_an.cell(row=tot, column=2, value=f"=SUM(B{first}:B{lastcat})")
        ws_an.cell(row=tot, column=3, value=f"=SUM(C{first}:C{lastcat})")
        ws_an.cell(row=tot, column=4, value=f"=IFERROR(C{tot}/B{tot},0)")
        for r in range(first, tot + 1):
            ws_an.cell(row=r, column=5, value=f"=IFERROR(C{r}/C${tot},0)")
            ws_an.cell(row=r, column=6, value=f"=(D{r}-D${tot})*100")
            for c in range(1, 7):
                cell = ws_an.cell(row=r, column=c)
                cell.font = f_bold if r == tot else f_body
                cell.border = box
                if r == tot:
                    cell.fill = fill_light
            ws_an.cell(row=r, column=2).number_format = "#,##0"
            ws_an.cell(row=r, column=3).number_format = "#,##0"
            ws_an.cell(row=r, column=4).number_format = "0.0%"
            ws_an.cell(row=r, column=5).number_format = "0.0%"
            ws_an.cell(row=r, column=6).number_format = '+0.0;-0.0;0.0'
        ws_an.conditional_formatting.add(
            f"D{first}:D{lastcat}",
            ColorScaleRule(start_type="min", start_color="C6EFCE", mid_type="percentile", mid_value=50,
                           mid_color="FFEB9C", end_type="max", end_color="F8CBAD"))
        tables[name] = dict(first=first, last=lastcat, total=tot, hdr=hdr)
        row = tot + 3

    # Cross-tab: Geography x Gender churn rate
    geos, genders = sorted(df["Geography"].unique()), sorted(df["Gender"].unique())
    ws_an.cell(row=row, column=1, value="Churn rate: Geography x Gender").font = Font(name=FONT, size=12, bold=True, color=NAVY)
    hdr = row + 1
    ws_an.cell(row=hdr, column=1, value="Geography \\ Gender")
    for c, g in enumerate(genders, 2):
        ws_an.cell(row=hdr, column=c, value=g)
    style_header(ws_an, hdr, 1, 1 + len(genders))
    for k, geo in enumerate(geos):
        r = hdr + 1 + k
        ws_an.cell(row=r, column=1, value=geo).font = f_body
        ws_an.cell(row=r, column=1).border = box
        for c in range(2, 2 + len(genders)):
            col = get_column_letter(c)
            cell = ws_an.cell(row=r, column=c, value=(
                f"=IFERROR(COUNTIFS({rng('Geography')},$A{r},{rng('Gender')},{col}${hdr},{EX},1)"
                f"/COUNTIFS({rng('Geography')},$A{r},{rng('Gender')},{col}${hdr}),0)"))
            cell.number_format, cell.font, cell.border = "0.0%", f_body, box
    ws_an.conditional_formatting.add(
        f"B{hdr+1}:{get_column_letter(1+len(genders))}{hdr+len(geos)}",
        ColorScaleRule(start_type="min", start_color="C6EFCE", end_type="max", end_color="F8CBAD"))
    ws_an.column_dimensions["A"].width = 24
    for col in "BCDEF":
        ws_an.column_dimensions[col].width = 18
    ws_an.freeze_panes = "A3"

    # ---------------- Dashboard ----------------
    ws_dash.sheet_view.showGridLines = False
    for c in range(1, 15):
        ws_dash.column_dimensions[get_column_letter(c)].width = 13
    ws_dash.column_dimensions["A"].width = 3
    ws_dash["B1"] = "Bank Customer Churn Dashboard"
    ws_dash["B1"].font = f_title
    ws_dash["B2"] = f"Data source: {source_label}  |  Records: "
    ws_dash["B2"].font = f_note
    ws_dash["H2"] = f"=COUNTA({rng('CustomerId')})"
    ws_dash["H2"].font, ws_dash["H2"].number_format = f_note, "#,##0"
    ws_dash["H2"].alignment = Alignment(horizontal="left")

    kpis = [
        ("Total Customers", f"=COUNTA({rng('CustomerId')})", "#,##0"),
        ("Churned Customers", f"=SUM({EX})", "#,##0"),
        ("Churn Rate", f"=IFERROR(SUM({EX})/COUNTA({rng('CustomerId')}),0)", "0.0%"),
        ("Avg Age: Churned", f"=IFERROR(AVERAGEIFS({rng('Age')},{EX},1),0)", "0.0"),
        ("Avg Age: Retained", f"=IFERROR(AVERAGEIFS({rng('Age')},{EX},0),0)", "0.0"),
        ("Balance Held by Churned", f"=SUMIFS({rng('Balance')},{EX},1)", "#,##0"),
    ]
    for i, (label, formula, fmt) in enumerate(kpis):
        c1 = 2 + i * 2
        ws_dash.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c1 + 1)
        ws_dash.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c1 + 1)
        a, b = ws_dash.cell(row=4, column=c1, value=label), ws_dash.cell(row=5, column=c1, value=formula)
        a.font, a.alignment = Font(name=FONT, size=9, bold=True, color="595959"), Alignment(horizontal="center")
        b.font, b.alignment, b.number_format = Font(name=FONT, size=16, bold=True, color=NAVY), Alignment(horizontal="center", vertical="center"), fmt
        for r in (4, 5):
            for c in (c1, c1 + 1):
                ws_dash.cell(row=r, column=c).fill = fill_light
        ws_dash.row_dimensions[5].height = 30

    def add_chart(dim, title, anchor, color):
        t = tables[dim]
        ch = BarChart()
        ch.type, ch.style, ch.title = "col", 10, title
        ch.add_data(Reference(ws_an, min_col=4, min_row=t["hdr"], max_row=t["last"]), titles_from_data=True)
        ch.set_categories(Reference(ws_an, min_col=1, min_row=t["first"], max_row=t["last"]))
        ch.legend = None
        ch.y_axis.numFmt, ch.y_axis.majorGridlines = "0%", None
        ch.y_axis.delete = ch.x_axis.delete = False
        ch.y_axis.scaling.min = 0
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.dataLabels.showSerName = ch.dataLabels.showCatName = ch.dataLabels.showLegendKey = False
        ch.series[0].graphicalProperties.solidFill = color
        ch.height, ch.width = 7.2, 15.2
        ws_dash.add_chart(ch, anchor)

    charts = [("Geography", "Churn Rate by Geography", "B7", BLUE), ("Age Group", "Churn Rate by Age Group", "H7", BLUE),
              ("NumOfProducts", "Churn Rate by Number of Products", "B23", "C55A11"),
              ("Activity Status", "Churn Rate: Active vs Inactive", "H23", "C55A11"),
              ("Credit Band", "Churn Rate by Credit Band", "B39", BLUE), ("Gender", "Churn Rate by Gender", "H39", BLUE)]
    for dim, title, anchor, color in charts:
        add_chart(dim, title, anchor, color)

    # Auto insights block
    r0 = 56
    ws_dash.cell(row=r0, column=2, value="Auto-generated insights (update automatically with the data)").font = Font(name=FONT, size=12, bold=True, color=NAVY)
    heads = ["Dimension", "", "Highest-churn segment", "", "Churn rate", "Lowest-churn segment", "", "Churn rate"]
    for c, h in enumerate(heads, 2):
        ws_dash.cell(row=r0 + 1, column=c, value=h)
    style_header(ws_dash, r0 + 1, 2, 9)
    for a_, b_ in ((2, 3), (4, 5), (7, 8)):
        ws_dash.merge_cells(start_row=r0 + 1, start_column=a_, end_row=r0 + 1, end_column=b_)
    for k, (dim, _) in enumerate(dims):
        t, r = tables[dim], r0 + 2 + k
        for a_, b_ in ((2, 3), (4, 5), (7, 8)):
            ws_dash.merge_cells(start_row=r, start_column=a_, end_row=r, end_column=b_)
        lab, rate = f"Analysis!$A${t['first']}:$A${t['last']}", f"Analysis!$D${t['first']}:$D${t['last']}"
        ws_dash.cell(row=r, column=2, value=dim)
        ws_dash.cell(row=r, column=4, value=f"=INDEX({lab},MATCH(MAX({rate}),{rate},0))")
        ws_dash.cell(row=r, column=6, value=f"=MAX({rate})")
        ws_dash.cell(row=r, column=7, value=f"=INDEX({lab},MATCH(MIN({rate}),{rate},0))")
        ws_dash.cell(row=r, column=9, value=f"=MIN({rate})")
        for c in range(2, 10):
            cell = ws_dash.cell(row=r, column=c)
            cell.font, cell.border = f_body, box
        ws_dash.cell(row=r, column=6).number_format = ws_dash.cell(row=r, column=9).number_format = "0.0%"
        for c in (4, 7):
            ws_dash.cell(row=r, column=c).alignment = Alignment(horizontal="left")

    # ---------------- README sheet ----------------
    ws_readme.sheet_view.showGridLines = False
    ws_readme.column_dimensions["A"].width = 3
    ws_readme.column_dimensions["B"].width = 26
    ws_readme.column_dimensions["C"].width = 95
    ws_readme["B1"] = "Bank Customer Churn Analysis (Excel)"
    ws_readme["B1"].font = f_title
    lines = [
        ("Data source", source_label),
        ("Records", f"{n:,} customers, one row per customer"),
        ("Dashboard", "KPI cards, 6 churn-rate charts and an auto-generated list of the highest/lowest churn segments."),
        ("Analysis", "One table per dimension: customers, churned, churn rate, share of all churners, and gap vs overall rate. Plus a Geography x Gender cross-tab."),
        ("Data", "Raw columns A to N from the CSV, plus 5 helper columns (formulas) that group customers into bands."),
        ("", ""),
        ("Helper column rules (assumptions)", ""),
        ("Age Group", "<30, 30-39, 40-49, 50-59, 60+"),
        ("Credit Band", "FICO-style bands: Poor <580, Fair 580-669, Good 670-739, Very Good 740-799, Excellent 800+"),
        ("Balance Segment", "Zero, up to 100K, 100K to 150K, above 150K (dataset currency is not stated; thresholds are my choice)"),
        ("Tenure Band", "0-2, 3-5, 6-8, 9-10 years"),
        ("Activity Status", "IsActiveMember 1 = Active, 0 = Inactive"),
        ("Churn definition", "Exited = 1 means the customer left the bank. Churn rate = churned / customers."),
        ("", ""),
        ("Refresh with new data", "Run: python src/build_workbook.py --input <your csv>  (CSV needs the same 14 columns). Or paste new rows into Data and extend the formula ranges."),
    ]
    for i, (a, b) in enumerate(lines, start=3):
        ws_readme.cell(row=i, column=2, value=a).font = f_bold
        cell = ws_readme.cell(row=i, column=3, value=b)
        cell.font, cell.alignment = f_body, Alignment(wrap_text=True, vertical="top")
    ws_readme["B9"].font = Font(name=FONT, size=11, bold=True, color=NAVY)

    wb.active = 1
    wb.save(out)
    print(f"Saved {out}  ({n:,} rows, source: {source_label})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--input")
    ap.add_argument("--output", default=str(ROOT / "excel" / "Bank_Churn_Analysis.xlsx"))
    a = ap.parse_args()
    if a.input:
        src, label = Path(a.input), f"Custom file: {Path(a.input).name}"
    elif REAL.exists():
        src, label = REAL, "Kaggle: Churn_Modelling.csv (Bank Customer Churn)"
    else:
        src, label = SAMPLE, "SYNTHETIC SAMPLE (demo only) - add the Kaggle file for real results"
        print("NOTE: data/Churn_Modelling.csv not found, using the synthetic sample.")
    build(load(src), a.output, label)
